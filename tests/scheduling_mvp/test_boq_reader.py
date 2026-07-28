from __future__ import annotations

from io import BytesIO

import openpyxl
import pytest

from construction_ai_scheduling.domain.errors import BOQImportError
from construction_ai_scheduling.infrastructure.boq_reader import TabularBOQReader


def test_csv_reader_suggests_unambiguous_columns():
    content = b"Item No,Description,Qty,UOM\nA-1,Concrete,10,m3\n"
    preview = TabularBOQReader().read("boq.csv", content)
    mapping = TabularBOQReader.suggest_mapping(preview.columns)

    assert mapping == {
        "item_code": "Item No",
        "description": "Description",
        "quantity": "Qty",
        "unit": "UOM",
    }
    rows = TabularBOQReader.extract_rows(preview.frame, mapping)
    assert rows[0]["quantity"] == 10.0
    assert rows[0]["source_row_number"] == 2


def test_reader_does_not_guess_ambiguous_description_column():
    content = b"Description,Item Description,Qty,UOM\nA,B,10,m3\n"
    preview = TabularBOQReader().read("boq.csv", content)
    assert TabularBOQReader.suggest_mapping(preview.columns)["description"] is None


def test_reader_preserves_invalid_quantity_for_review():
    content = b"Description,Qty,UOM\nConcrete,TBC,m3\n"
    preview = TabularBOQReader().read("boq.csv", content)
    rows = TabularBOQReader.extract_rows(preview.frame, {
        "item_code": None,
        "description": "Description",
        "quantity": "Qty",
        "unit": "UOM",
    })
    assert rows[0]["quantity"] is None
    assert rows[0]["quantity_raw"] == "TBC"


def test_xlsx_reader_selects_requested_worksheet():
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "Cover"
    first.append(["Notes"])
    first.append(["Demonstration only"])
    boq = workbook.create_sheet("BOQ")
    boq.append(["Description", "Quantity", "Unit"])
    boq.append(["Concrete", 12, "m3"])
    stream = BytesIO()
    workbook.save(stream)

    preview = TabularBOQReader().read("boq.xlsx", stream.getvalue(), sheet_name="BOQ")
    assert preview.sheet_names == ("Cover", "BOQ")
    assert preview.selected_sheet == "BOQ"
    assert preview.frame.iloc[0]["Quantity"] == 12


def test_xlsx_reader_detects_arabic_header_after_introductory_rows():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "جدول الكميات"
    worksheet.append(["جدول الكميات للمشروع"])
    worksheet.append(["مشروع تجريبي - بيانات توضيحية"])
    worksheet.append([])
    worksheet.append(["رقم البند", "الوصف", "الكمية", "الوحدة"])
    worksheet.append(["A-01", "أعمال خرسانة مسلحة", 125, "م³"])
    stream = BytesIO()
    workbook.save(stream)

    preview = TabularBOQReader().read("arabic-boq.xlsx", stream.getvalue())
    mapping = TabularBOQReader.suggest_mapping(preview.columns)
    rows = TabularBOQReader.extract_rows(preview.frame, mapping)

    assert preview.header_row_number == 4
    assert preview.columns == ("رقم البند", "الوصف", "الكمية", "الوحدة")
    assert len(rows) == 1
    assert rows[0]["source_row_number"] == 5
    assert rows[0]["description"] == "أعمال خرسانة مسلحة"
    assert rows[0]["unit"] == "م³"


def test_xlsx_reader_allows_explicit_header_selection_without_dropping_data_rows():
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Description", "Quantity", "Unit"])
    worksheet.append(["Introductory legend", "Not BOQ data", "Notes"])
    worksheet.append(["Project title"])
    worksheet.append(["Description", "Quantity", "Unit"])
    worksheet.append(["Concrete", 10, "m3"])
    worksheet.append(["بند يحتاج مراجعة", None, "عدد"])
    stream = BytesIO()
    workbook.save(stream)

    reader = TabularBOQReader()
    automatic = reader.read("boq-with-intro.xlsx", stream.getvalue())
    selected = reader.read("boq-with-intro.xlsx", stream.getvalue(), header_row=4)

    assert automatic.header_row_number == 1
    assert 4 in dict(automatic.header_row_options)
    assert list(selected.frame.index) == [5, 6]
    assert selected.frame["Description"].tolist() == ["Concrete", "بند يحتاج مراجعة"]


def test_reader_rejects_unsupported_or_empty_files():
    reader = TabularBOQReader()
    with pytest.raises(BOQImportError, match="CSV or XLSX"):
        reader.read("boq.xls", b"data")
    with pytest.raises(BOQImportError, match="empty"):
        reader.read("boq.csv", b"")
